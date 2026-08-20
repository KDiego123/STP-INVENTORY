"""Actualiza descripción, marca y modelo sin reconstruir inventario.

El archivo maestro contiene hojas de auditoría con el estado anterior y el
resultado normalizado. Este cargador solo modifica los códigos auditados,
conserva ``inventario.id`` y rechaza la operación si detecta una edición
posterior que no coincide ni con el estado original ni con el resultado final.
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from importar_inventario_maestro import database_config, load_env


EXPECTED_MODELS = 169
EXPECTED_BRANDS = 65
EXPECTED_OVERLAP = 2
EXPECTED_CHANGES = 232
ADVISORY_LOCK = 2026082001


@dataclass
class AttributeChange:
    code: str
    original_description: str
    final_description: str
    update_model: bool = False
    final_model: str | None = None
    update_brand: bool = False
    original_brand: str | None = None
    final_brand: str | None = None


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split())


def optional_text(value: Any) -> str | None:
    text = clean_text(value)
    return text or None


def sheet_rows(worksheet: Any) -> list[tuple[Any, ...]]:
    return [
        row
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in row)
    ]


def read_changes(path: Path) -> tuple[dict[str, AttributeChange], int]:
    if not path.is_file():
        raise RuntimeError(f"No existe el maestro: {path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    required = {
        "INVENTARIO MAESTRO",
        "CONTROL EXTRACCION MODELOS",
        "CONTROL EXTRACCION MARCAS",
    }
    missing = sorted(required - set(workbook.sheetnames))
    if missing:
        raise RuntimeError("Faltan hojas obligatorias: " + ", ".join(missing))

    master_rows = sheet_rows(workbook["INVENTARIO MAESTRO"])
    master: dict[str, tuple[str, str | None, str | None]] = {}
    for row in master_rows:
        code = clean_text(row[1])
        if not code or code in master:
            raise RuntimeError(f"Código vacío o duplicado en el maestro: {code!r}")
        master[code] = (clean_text(row[3]), optional_text(row[8]), optional_text(row[9]))

    model_rows = sheet_rows(workbook["CONTROL EXTRACCION MODELOS"])
    brand_rows = sheet_rows(workbook["CONTROL EXTRACCION MARCAS"])
    if len(model_rows) != EXPECTED_MODELS:
        raise RuntimeError(
            f"Se esperaban {EXPECTED_MODELS} auditorías de modelo y se encontraron {len(model_rows)}."
        )
    if len(brand_rows) != EXPECTED_BRANDS:
        raise RuntimeError(
            f"Se esperaban {EXPECTED_BRANDS} auditorías de marca y se encontraron {len(brand_rows)}."
        )

    changes: dict[str, AttributeChange] = {}
    model_codes: set[str] = set()
    brand_codes: set[str] = set()

    for row in model_rows:
        code = clean_text(row[2])
        if code in model_codes:
            raise RuntimeError(f"Código duplicado en auditoría de modelos: {code}")
        model_codes.add(code)
        if code not in master:
            raise RuntimeError(f"El código auditado no existe en el maestro: {code}")
        original_description = clean_text(row[3])
        final_description, _, final_model = master[code]
        audit_model = optional_text(row[5])
        if final_model != audit_model:
            raise RuntimeError(f"El modelo final no coincide con la auditoría: {code}")
        changes[code] = AttributeChange(
            code=code,
            original_description=original_description,
            final_description=final_description,
            update_model=True,
            final_model=final_model,
        )

    for row in brand_rows:
        code = clean_text(row[2])
        if code in brand_codes:
            raise RuntimeError(f"Código duplicado en auditoría de marcas: {code}")
        brand_codes.add(code)
        if code not in master:
            raise RuntimeError(f"El código auditado no existe en el maestro: {code}")
        original_description = clean_text(row[3])
        final_description, final_brand, _ = master[code]
        audit_brand = optional_text(row[6])
        if final_brand != audit_brand:
            raise RuntimeError(f"La marca final no coincide con la auditoría: {code}")
        change = changes.get(code)
        if change is None:
            change = AttributeChange(
                code=code,
                original_description=original_description,
                final_description=final_description,
            )
            changes[code] = change
        elif change.original_description != original_description:
            raise RuntimeError(f"Las auditorías no coinciden en la descripción original: {code}")
        change.final_description = final_description
        change.update_brand = True
        change.original_brand = optional_text(row[5])
        change.final_brand = final_brand

    overlap = len(model_codes & brand_codes)
    if overlap != EXPECTED_OVERLAP or len(changes) != EXPECTED_CHANGES:
        raise RuntimeError(
            "La unión de auditorías no coincide con la versión aprobada: "
            f"intersección={overlap}, códigos={len(changes)}."
        )
    return changes, len(master)


def verify_schema(cursor: Any) -> None:
    cursor.execute(
        """
        SELECT column_name, udt_name, character_maximum_length
        FROM information_schema.columns
        WHERE table_schema = 'public'
          AND table_name = 'inventario'
          AND column_name IN (
              'id', 'codigo', 'descripcion', 'marca', 'modelo', 'actualizado_en'
          )
        """
    )
    columns = {name: (kind, length) for name, kind, length in cursor.fetchall()}
    required = {
        "id": "int8",
        "codigo": "varchar",
        "descripcion": "text",
        "marca": "varchar",
        "modelo": "varchar",
        "actualizado_en": "timestamptz",
    }
    missing = [name for name, kind in required.items() if columns.get(name, (None,))[0] != kind]
    if missing:
        raise RuntimeError(
            "La estructura de public.inventario no coincide con backupinventario.sql: "
            + ", ".join(missing)
        )
    for name in ("marca", "modelo"):
        length = columns[name][1]
        if length is not None and length < 100:
            raise RuntimeError(f"La columna {name} admite menos de 100 caracteres.")


def same_optional(left: Any, right: Any) -> bool:
    return optional_text(left) == optional_text(right)


def validate_current(
    changes: dict[str, AttributeChange],
    current: dict[str, tuple[str, str | None, str | None]],
) -> list[str]:
    conflicts: list[str] = []
    for code, change in changes.items():
        if code not in current:
            conflicts.append(f"{code}: no existe en public.inventario")
            continue
        description, brand, model = current[code]
        if description not in {change.original_description, change.final_description}:
            conflicts.append(f"{code}: la descripción fue modificada después de la importación")
        if change.update_brand and not (
            same_optional(brand, change.original_brand)
            or same_optional(brand, change.final_brand)
        ):
            conflicts.append(
                f"{code}: marca actual {brand!r}; se esperaba {change.original_brand!r} "
                f"o {change.final_brand!r}"
            )
        if change.update_model and not (
            same_optional(model, None) or same_optional(model, change.final_model)
        ):
            conflicts.append(
                f"{code}: modelo actual {model!r}; se esperaba vacío o {change.final_model!r}"
            )
    return conflicts


def update_database(
    changes: dict[str, AttributeChange],
    master_count: int,
    env_path: Path,
    apply: bool,
) -> int:
    load_env(env_path)
    try:
        import psycopg
    except ImportError as exc:
        raise RuntimeError("Falta psycopg. Ejecute pip install -r requirements.txt.") from exc

    updated = 0
    with psycopg.connect(**database_config()) as connection:
        with connection.transaction():
            with connection.cursor() as cursor:
                cursor.execute("SELECT pg_advisory_xact_lock(%s)", (ADVISORY_LOCK,))
                verify_schema(cursor)
                cursor.execute("LOCK TABLE public.inventario IN SHARE ROW EXCLUSIVE MODE")
                cursor.execute("SELECT count(*) FROM public.inventario")
                database_count = int(cursor.fetchone()[0])
                codes = list(changes)
                cursor.execute(
                    """
                    SELECT codigo, descripcion, marca, modelo
                    FROM public.inventario
                    WHERE codigo = ANY(%s)
                    """,
                    (codes,),
                )
                current = {
                    code: (description, brand, model)
                    for code, description, brand, model in cursor.fetchall()
                }
                conflicts = validate_current(changes, current)
                if conflicts:
                    preview = "\n".join(f"  - {message}" for message in conflicts[:30])
                    extra = "" if len(conflicts) <= 30 else f"\n  ... y {len(conflicts) - 30} más"
                    raise RuntimeError(
                        f"Se detectaron {len(conflicts)} conflictos; no se modificó la base:\n"
                        + preview
                        + extra
                    )

                pending: list[tuple[AttributeChange, str | None, str | None]] = []
                for code, change in changes.items():
                    description, brand, model = current[code]
                    final_brand = change.final_brand if change.update_brand else brand
                    final_model = change.final_model if change.update_model else model
                    if (
                        description != change.final_description
                        or not same_optional(brand, final_brand)
                        or not same_optional(model, final_model)
                    ):
                        pending.append((change, final_brand, final_model))

                print(f"Artículos en la base: {database_count}")
                print(f"Artículos en el maestro: {master_count}")
                print(f"Códigos auditados: {len(changes)}")
                print(f"Códigos pendientes de actualización: {len(pending)}")

                if not apply:
                    print("Validación de BD completada; no se realizaron cambios.")
                    return len(pending)

                for change, final_brand, final_model in pending:
                    cursor.execute(
                        """
                        UPDATE public.inventario
                        SET descripcion = %s,
                            marca = %s,
                            modelo = %s,
                            actualizado_en = CURRENT_TIMESTAMP
                        WHERE codigo = %s
                        """,
                        (
                            change.final_description,
                            final_brand,
                            final_model,
                            change.code,
                        ),
                    )
                    if cursor.rowcount != 1:
                        raise RuntimeError(f"No se pudo actualizar de forma única: {change.code}")
                    updated += 1

                cursor.execute(
                    """
                    SELECT codigo, descripcion, marca, modelo
                    FROM public.inventario
                    WHERE codigo = ANY(%s)
                    """,
                    (codes,),
                )
                final = {
                    code: (description, brand, model)
                    for code, description, brand, model in cursor.fetchall()
                }
                remaining = []
                for code, change in changes.items():
                    description, brand, model = final[code]
                    if description != change.final_description:
                        remaining.append(code)
                    elif change.update_brand and not same_optional(brand, change.final_brand):
                        remaining.append(code)
                    elif change.update_model and not same_optional(model, change.final_model):
                        remaining.append(code)
                if remaining:
                    raise RuntimeError(
                        "La verificación posterior falló para: " + ", ".join(remaining[:30])
                    )
    return updated


def parse_args() -> argparse.Namespace:
    base = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(
        description="Actualiza descripción, marca y modelo sobre una base existente."
    )
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--validar-bd", action="store_true", help="Comprueba sin modificar.")
    mode.add_argument("--aplicar", action="store_true", help="Aplica en una transacción.")
    parser.add_argument("--archivo", type=Path, default=base / "INVENTARIO_MAESTRO_STP.xlsx")
    parser.add_argument("--env", type=Path, default=base / ".env")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        changes, master_count = read_changes(args.archivo)
        updated = update_database(changes, master_count, args.env, args.aplicar)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    if args.aplicar:
        print(f"Actualización confirmada: {updated} artículos modificados.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
